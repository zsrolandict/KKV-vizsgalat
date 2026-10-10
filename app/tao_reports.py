from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .tao_models import FAMILY_LABELS, CONTROL_LABELS, PE_LABELS
from .tao_engine import RESULT_LABELS
from .tao_graph import graph_image


ANSWER_LABELS = {'yes': 'Igen', 'no': 'Nem', 'unknown': 'Tisztázandó'}
MEMBERSHIP_LABELS = {'member': 'Tag / részvényes', 'not_member': 'Nem tag / részvényes', 'unknown': 'Tisztázandó'}


def evidence_text(evidence):
    source = evidence.get('source', '')
    if evidence.get('page'):
        source += f' · {evidence["page"]}. oldal'
    if evidence.get('quote'):
        source += '\nIdézet: ' + evidence['quote']
    return source


def text_cell(value):
    # Force literal strings, including user text starting with '=' or '+'.
    return str(value if value is not None else '')


def xlsx_report(data, calc, meta):
    wb = Workbook()
    matrix = wb.active
    matrix.title = 'Mátrix'
    matrix.append(['Vizsgálat', data.title])
    matrix.append(['Vizsgálati nap', data.as_of.isoformat()])
    matrix.append(['Állapot', meta['label']])
    matrix.append(['Értékelés', 'Szakértői döntések; a szavazati és irányítási jelzések nem végleges jogi minősítések.'])
    matrix.append([])
    matrix.append(['Vállalkozás', *[c.name for c in data.companies]])
    lookup = {frozenset((r['first'], r['second'])): r for r in calc['rows']}
    for a in data.companies:
        matrix.append([a.name, *['—' if a.id == b.id else lookup[frozenset((a.id, b.id))]['label'] for b in data.companies]])
    matrix.freeze_panes = 'B7'
    details = wb.create_sheet('Indokolások')
    details.append(['Cég A', 'Cég B', 'Jogi eredmény', 'Munkafolyamat', 'Jogalap', 'Indok', 'Forrás / oldal', 'Feltételezés', 'Hiányzó adat', 'Kapcsoltsági jelzés'])
    for row in calc['rows']:
        details.append([row['first_name'], row['second_name'], row['label'], row['stage_label'], row['basis'], row['reason'],
                        evidence_text(row['evidence']), '\n'.join(row['assumptions']), '\n'.join(row['missing']), '\n'.join(row['signals'])])
    details.freeze_panes = 'C2'
    details.auto_filter.ref = details.dimensions
    computations = wb.create_sheet('Befolyásszámítás')
    computations.append(['Cég A', 'Cég B', 'Jogosult', 'Célcég', 'Befolyás', 'Többség igazolt', 'Joghely', 'Tényazonosítók'])
    actor_names = {x.id: x.name for x in [*data.companies, *data.persons]}
    for row in calc['rows']:
        for value in row.get('influence_calculations', []):
            computations.append([row['first_name'], row['second_name'], value.get('owner_name') or actor_names[value['owner']],
                                 actor_names[value['company']], value['value'], 'Igen' if value['majority'] else 'Nem',
                                 value['basis'], ', '.join(value['fact_ids'])])
    facts = wb.create_sheet('Szavazati források')
    names = {x.id: x.name for x in [*data.companies, *data.persons]}
    facts.append(['Tulajdonos', 'Célcég', 'Tőke (%)', 'Szavazati mód', 'Szavazat (%) / feltétel', 'Kezdet', 'Vége (kizáró)', 'Ellenőrzött nap', 'Forrás', 'Indok', 'Tényazonosító', 'Részesedés jogállása', 'BVK-hozzárendelés ellenőrizve'])
    for fact in data.voting_facts:
        vote = '>50%' if fact.vote_bound == 'over_half' else fact.votes
        facts.append([names[fact.owner], names[fact.company], fact.capital, fact.vote_mode, vote,
                      fact.valid_from, fact.valid_to, fact.reviewed_as_of, fact.evidence.source, fact.reason, fact.id, fact.capacity, str(fact.attribution_reviewed)])
    relatives = wb.create_sheet('Rokonsági források')
    relatives.append(['Első személy', 'Második személy', 'Viszony', 'Kezdet', 'Vége (kizáró)', 'Ellenőrzött nap', 'Igazolt', 'Forrás / oldal', 'Tényazonosító'])
    for fact in data.family_facts:
        relatives.append([names[fact.first], names[fact.second], FAMILY_LABELS[fact.relationship],
                          fact.valid_from, fact.valid_to, fact.reviewed_as_of,
                          'Igen' if fact.confirmed else 'Nem', evidence_text(fact.evidence.model_dump(mode='json')), fact.id])
    rights = wb.create_sheet('Irányítási jogok')
    rights.append(['Jogosult', 'Célcég', 'Jogcím', 'Tagi jogállás', 'Feltétel fennáll', 'Együttes szavazat (%)', 'Igazolt', 'Indok', 'Forrás / oldal', 'Kezdet', 'Vége (kizáró)', 'Ellenőrzött nap', 'Tényazonosító'])
    for fact in data.control_facts:
        rights.append([names[fact.owner], names[fact.company], CONTROL_LABELS[fact.kind], MEMBERSHIP_LABELS[fact.membership],
                       ANSWER_LABELS[fact.condition], '>50%' if fact.aligned_bound == 'over_half' else fact.aligned_votes, 'Igen' if fact.confirmed else 'Nem', fact.reason,
                       evidence_text(fact.evidence.model_dump(mode='json')), fact.valid_from, fact.valid_to,
                       fact.reviewed_as_of, fact.id])
    managers = wb.create_sheet('Ügyvezetési tények')
    managers.append(['Cég A', 'Cég B', 'Közös vezetők', 'Egyezőség', 'Üzleti döntő befolyás', 'Pénzügyi döntő befolyás', 'Igazolt', 'Indok', 'Forrás / oldal', 'Kezdet', 'Vége (kizáró)', 'Ellenőrzött nap', 'Tényazonosító'])
    for fact in data.management_facts:
        managers.append([names[fact.first], names[fact.second], ', '.join(names[m] for m in fact.managers),
                         ANSWER_LABELS[fact.common_management], ANSWER_LABELS[fact.business_control], ANSWER_LABELS[fact.financial_control],
                         'Igen' if fact.confirmed else 'Nem', fact.reason, evidence_text(fact.evidence.model_dump(mode='json')),
                         fact.valid_from, fact.valid_to, fact.reviewed_as_of, fact.id])
    routes = wb.create_sheet('Irányítási útvonalak')
    routes.append(['Cég A', 'Cég B', 'Jogosult', 'Célcég', 'Útvonal', 'Jogalap', 'Irányítási tények', 'Szavazati tények'])
    for row in calc['rows']:
        for route in row.get('control_calculations', []):
            routes.append([row['first_name'], row['second_name'], names[route['owner']], names[route['company']],
                           ' → '.join(names[n] for n in route['route']), route['basis'],
                           ', '.join(route['control_ids']), ', '.join(route['fact_ids'])])
    pe = wb.create_sheet('Tao-telephelyek')
    pe.append(['Fővállalkozás','Telephely','Jogviszony','Adójogi jogállás','Igazolt','Indok','Forrás','Kezdet','Vége','Ellenőrzött nap'])
    for f in data.establishment_facts:
        pe.append([names[f.principal],names[f.establishment],PE_LABELS[f.kind],ANSWER_LABELS[f.tax_status],str(f.confirmed),f.reason,evidence_text(f.evidence.model_dump(mode='json')),f.valid_from,f.valid_to,f.reviewed_as_of])
    trusts = wb.create_sheet('BVK-tények')
    trusts.append(['Kezelt vagyon','Vagyonkezelők','Vagyonrendelők','Kedvezményezettek','Érintett cégek','Jogok ellenőrizve','Indok','Forrás'])
    for f in data.trust_facts:
        trusts.append([f.name,', '.join(names[x] for x in f.trustees),', '.join(names[x] for x in f.settlors),', '.join(names[x] for x in f.beneficiaries),', '.join(names[x] for x in f.holdings),str(f.rights_reviewed),f.reason,evidence_text(f.evidence.model_dump(mode='json'))])
    profile = wb.create_sheet('Vizsgálati keret')
    for key, value in [('Ügyazonosító', meta['id']), ('Verzió', meta['version']), ('Állapot', meta['label']),
                       ('Vizsgálati cél', data.purpose), ('Vállalt vizsgálati kör', data.scope),
                       ('Feltételezések', data.assumptions), ('Jogi időállapot', data.law_date),
                       ('Történeti alkalmazhatóság', data.law_applicability), ('Jogi forrás', data.law_source), ('Munkafolyamat-verzió', calc['rule_version']),
                       ('Jóváhagyó', meta.get('approver') or ''), ('Jóváhagyás', meta.get('approved_at') or '')]:
        profile.append([key, value])
    for sheet in wb:
        for row in sheet:
            for cell in row:
                cell.value = text_cell(cell.value)
                cell.data_type = 's'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1C3658')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 26
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def word_report(data, calc, meta):
    from .reports import configure, add_table, summary_box, new_section
    from .opinion_support import market_paragraphs, assumptions, add_authorities, signoff
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc=Document();configure(doc,data,calc,meta,tao=True)
    p=doc.add_heading('állásfoglalás',0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph('– Tao. törvény szerinti kapcsoltsági vizsgálat –',style='Subtitle');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(data.title,style='Subtitle')
    doc.add_paragraph(f'Vizsgálati nap: {data.as_of} · {meta["version"]}. ügyverzió',style='Caption')
    conclusion=(f'A {data.as_of} napjára elvégzett vizsgálat {len(data.companies)} vállalkozás {calc["total_pairs"]} cégpárjára terjed ki. '
        f'A rögzített, megerősített döntések alapján {calc["counts"]["related"]} pár kapcsolt, '
        f'{calc["counts"]["not_related"]} pár nem kapcsolt; {calc["counts"]["undetermined"]} pár minősítése még nem dönthető el.')
    summary_box(doc,meta['label'],conclusion,not meta['approved'])
    if not meta['approved']:doc.add_paragraph('TERVEZET – szakértői jóváhagyás nélkül végleges állásfoglalásként nem használható.',style='Caption')
    elif not calc['complete']:doc.add_paragraph('RÉSZLEGES ÁLLÁSFOGLALÁS – a nem eldöntött cégpárok nem minősülnek nem kapcsoltnak.',style='Caption')
    doc.add_heading('1. A megbízott feladata és a vizsgálat kerete',1)
    doc.add_paragraph(f'A megbízás tárgya a felsorolt vállalkozások közötti kapcsoltság értékelése a társasági adóról és az osztalékadóról '
        f'szóló 1996. évi LXXXI. törvény (Tao.) 4. § 23. pontjának a rögzített időállapotban alkalmazandó rendelkezései alapján. '
        f'Megbízó: {data.client or "nincs külön megadva"}. A vizsgálat célja: {data.purpose}.')
    doc.add_paragraph('Vállalt vizsgálati kör: '+(data.scope or 'A vizsgálati kör és korlátai még nem kerültek külön meghatározásra.'))
    doc.add_paragraph(f'Jogi időállapot: {data.law_date or "ellenőrizendő"}. Forrás: {data.law_source or "nincs rögzítve"}. '
                      f'Történeti alkalmazhatóság: {data.law_applicability or "külön indok nincs rögzítve"}.')
    doc.add_heading('2. Vezetői megállapítások és alkalmazott módszer',1)
    doc.add_paragraph(conclusion)
    doc.add_paragraph('A vizsgálatban külön értékeljük a dokumentált tulajdoni és szavazati jogokat, a közvetett befolyási útvonalakat, '
        'az igazolt hozzátartozói összeszámítást, a meghatározó irányítási jogokat és az ügyvezetési tényeket. '
        'Tao-telephely vagy bizalmi vagyonkezelés esetén a hozzájuk tartozó külön tényállás és forrás is szükséges. '
        'A számított jelzést a megerősített szakértői jogi döntéstől elkülönítjük; a hiányzó adatot nem értelmezzük negatív minősítésként.')
    add_table(doc,['Cég A','Cég B','Jogi eredmény'],[(r['first_name'],r['second_name'],r['label']) for r in calc['rows']],[6,6,5.5])
    doc.add_heading('3. Cégpáronkénti tényállás, jogi értékelés és következtetés',1)
    names={a.id:a.name for a in [*data.companies,*data.persons]}
    for row in calc['rows']:
        doc.add_heading(row['first_name']+' – '+row['second_name'],2)
        p=doc.add_paragraph();p.add_run('Megállapítás: ').bold=True;p.add_run(row['label']+'. '+row['stage_label']+'.')
        if row['result']=='undetermined':doc.add_paragraph('A cégpár végleges kapcsoltsági minősítése még nem igazolható. '
            'Az alábbi tények és jelzések további szakértői értékelést vagy bizonyítást igényelnek.')
        doc.add_paragraph('Szakértői indokolás: '+row['reason'])
        doc.add_paragraph('Alkalmazott jogalap: '+(row['basis'] or 'A konkrét rendelkezés még nincs rögzítve.'))
        if row['evidence'].get('source'):doc.add_paragraph('Igazoló forrás: '+evidence_text(row['evidence']))
        for signal in row['signals']:doc.add_paragraph('A rögzített tényekből származó jelzés: '+signal)
        for v in row.get('influence_calculations',[]):
            doc.add_paragraph(f'Befolyásszámítás: {v.get("owner_name") or names[v["owner"]]} → {names[v["company"]]}: {v["value"]}. '
                f'Az alkalmazott számítás jogalapja: {v["basis"]}. Felhasznált szavazati tények: {", ".join(v["fact_ids"])}.')
        for f in row.get('family_facts',[]):doc.add_paragraph('Igazolt rokonság: '+names[f['first']]+' – '+names[f['second']]+': '+FAMILY_LABELS[f['relationship']]+'. Forrás: '+evidence_text(f['evidence']))
        for route in row.get('control_calculations',[]):doc.add_paragraph('Meghatározó befolyás útvonala: '+' → '.join(names[n] for n in route['route'])+'. Jogalap: '+route['basis']+'.')
        for f in row.get('control_facts',[]):
            doc.add_paragraph('Irányítási tény: '+names[f['owner']]+' → '+names[f['company']]+'. '+CONTROL_LABELS[f['kind']]+'. '
                'Tagi jogállás: '+MEMBERSHIP_LABELS[f['membership']]+'. Feltétel fennállása: '+ANSWER_LABELS[f['condition']]+'. '+f['reason']+'. Forrás: '+evidence_text(f['evidence']))
            if f['kind']=='voting_agreement':doc.add_paragraph('Megállapodás szerinti együttes szavazat: '+('>50%' if f['aligned_bound']=='over_half' else str(f['aligned_votes'])+'%' if f['aligned_votes'] is not None else 'tisztázandó')+'.')
        for f in row.get('management_facts',[]):doc.add_paragraph('Ügyvezetési tény: '+', '.join(names[m] for m in f['managers'])+'. '
            'Egyezőség: '+ANSWER_LABELS[f['common_management']]+', üzleti döntő befolyás: '+ANSWER_LABELS[f['business_control']]+', '
            'pénzügyi döntő befolyás: '+ANSWER_LABELS[f['financial_control']]+'. '+f['reason']+'. Forrás: '+evidence_text(f['evidence']))
        for f in row.get('establishment_facts',[]):doc.add_paragraph('Telephelyi tény: '+names[f['principal']]+' → '+names[f['establishment']]+'. '+PE_LABELS[f['kind']]+'. '+f['reason']+'. Forrás: '+evidence_text(f['evidence']))
        for f in row.get('trust_facts',[]):doc.add_paragraph('Bizalmi vagyonkezelési tényállás: '+f['name']+'. '+f['reason']+'. Forrás: '+evidence_text(f['evidence']))
        for label,values in [('Ügyre alkalmazott feltételezés',row['assumptions']),('Tisztázandó tény / következő lépés',row['missing'])]:
            for v in values:doc.add_paragraph(label+': '+v)
    doc.add_heading('4. Piaci összefüggések és a vizsgálatok elhatárolása',1)
    for text in market_paragraphs(data,tao=True):doc.add_paragraph(text)
    doc.add_heading('5. Feltételezések, korlátozások és nyitott kérdések',1)
    for text in assumptions(data,tao=True):doc.add_paragraph(text,style='List Bullet')
    pending=[r for r in calc['rows'] if r['result']=='undetermined' or not r['confirmed']]
    for row in pending:doc.add_paragraph(row['first_name']+' – '+row['second_name']+': '+(' '.join(row['missing']) or 'Indokolt, forrással igazolt minősítés és megerősítés szükséges.'),style='List Bullet')
    if not pending:doc.add_paragraph('Minden cégpárra megerősített minősítés áll rendelkezésre.')
    doc.add_heading('6. Alkalmazott jogszabályok és bírósági döntések',1)
    doc.add_paragraph('Tao.: 1996. évi LXXXI. törvény, a rögzített kapcsoltsági döntésekben megjelölt rendelkezések. '
                      'A Ptk. szerinti szavazati és befolyási számítások joghelyei a cégpáronkénti indokolásban szerepelnek. '
                      'Az alkalmazandó időállapotot és az ügyre való alkalmazhatóságot a szakértő ellenőrzi.')
    add_authorities(doc,data)
    concluding_start=len(doc.paragraphs)
    doc.add_heading('7. Összefoglalás és szakértői jóváhagyás',1)
    summary_box(doc,'Megállapítás',conclusion,not meta['approved'])
    if meta['approved']:
        stamp=datetime.fromisoformat(meta['approved_at']).astimezone(ZoneInfo('Europe/Budapest')).strftime('%Y.%m.%d. %H:%M')
        doc.add_paragraph(f'Jóváhagyó: {meta["approver"]}. Jóváhagyás: {stamp} (Europe/Budapest).')
        if meta.get('note'):doc.add_paragraph('Jóváhagyási korlátok / megjegyzés: '+meta['note'])
    else:doc.add_paragraph('Szakértői jóváhagyás még nem történt.')
    doc.add_paragraph('A KKV-minősítés és az ügyletszintű transzferár-kötelezettség külön vizsgálat. '
                      'A kapcsoltság megállapítása önmagában nem igazolja valamennyi nyilvántartási feltétel fennállását. '
                      'A rendszerbeli jóváhagyás nem elektronikus aláírás.')
    signoff(doc,data,meta)
    for p in doc.paragraphs[concluding_start:-1]:p.paragraph_format.keep_with_next=True
    for row in doc.tables[-1].rows:
        for cell in row.cells:
            for p in cell.paragraphs:p.paragraph_format.keep_with_next=True
    new_section(doc,landscape=True)
    doc.add_heading('1. melléklet – cégháló és rögzített kapcsolatok',1)
    doc.add_paragraph('A mentett elrendezésű ábra a vizsgálati napon fennálló rögzített tényeket mutatja. '
        'A * jel a tulajdon és szavazat azonosságának feltételezését jelöli. A nyilak nem végleges jogi minősítések.',style='Caption')
    png=graph_image(data)
    from PIL import Image
    with Image.open(BytesIO(png)) as image:ratio=image.width/image.height
    doc.add_picture(BytesIO(png),width=Cm(min(25,12.5*ratio)))
    output=BytesIO();doc.save(output);return output.getvalue()
