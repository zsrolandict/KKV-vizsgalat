"""Case-specific market analysis and reviewed legal authorities for both opinion types."""
from .report_text import hu


def market_paragraphs(data, tao=False):
    paragraphs=[]
    if data.market_analysis: paragraphs.append(data.market_analysis)
    if not tao:
        names={a.id:a.name for a in [*data.companies,*data.persons]}
        for d in data.decisions:
            if d.basis=='persons' or d.market or d.acting_together:
                paragraphs.append(f'{names[d.first]} – {names[d.second]}: '+
                    f'közös fellépés: {d.acting_together or "nincs dokumentálva"}. '+
                    f'Azonos / szomszédos piac értékelése: {d.market or "nincs dokumentálva"}. '+
                    f'Forrás: {d.source or "nincs megadva"}. '+
                    (('Megerősített vizsgálati évek: '+', '.join(map(str,d.confirmed_years if d.confirmed_years is not None else data.years))+'.') if d.confirmed else 'A kapcsolati döntés ellenőrizendő.'))
        if paragraphs or data.families:
            paragraphs.append('Természetes személyeken keresztüli kapcsolódásnál a közös fellépés és az azonos vagy szomszédos piac '
                              'ügyre szabott vizsgálata szükséges. A rokonság, az azonos TEÁOR-kód vagy a puszta 100%-os '
                              'érzékenységvizsgálat önmagában nem helyettesíti ezt az értékelést.')
    elif data.market_analysis:
        paragraphs.append('A piaci összefüggés a Tao szerinti tényállás kiegészítése. A Kkv. törvény természetes személyi '
                          'kapcsolódásának piacfeltételét nem alkalmazzuk automatikusan a Tao. 4. § 23. pontjának valamennyi jogalapjára.')
    if not paragraphs:
        paragraphs.append('Külön piaci értékelés nincs rögzítve. A dokumentum az azonos vagy szomszédos piac fennállásáról '
                          'nem tesz önálló megállapítást; ha ez a választott jogalaphoz szükséges, ügyre szabott tényállással és forrással kell kiegészíteni.')
    return paragraphs


def assumptions(data, tao=False):
    lines=[line.strip() for line in data.assumptions.splitlines() if line.strip()]
    if tao:
        names={a.id:a.name for a in [*data.companies,*data.persons]}
        from .tao_engine import active
        for f in data.voting_facts:
            if f.vote_mode=='ownership_default' and active(f,data.as_of):
                lines.append(f'{names[f.owner]} → {names[f.company]}: a {hu(f.capital)}%-os tulajdoni arányt '
                             'szavazati arányként kezeljük. Ez rögzített munkafeltételezés; eltérő szavazati jog igazolása újraszámítást igényel.')
        lines.extend(d.assumptions for d in data.decisions if d.as_of==data.as_of and d.assumptions)
    else:
        if any(f.estimated for f in data.financials):
            lines.append('A becsült / évesített beszámolóadatok a megadott módszer és forrás alapján szerepelnek; '
                         'a végleges beszámoló rendelkezésre állásakor a számítást ismételten el kell végezni.')
    if not lines:lines.append('Külön, ügyre szabott feltételezés nincs rögzítve. Nem feltételezzük automatikusan '
                             'más irányítási szerződés, eltérő szavazati jog vagy közjogi részesedés hiányát.')
    lines.append('A megállapítás a rögzített adatokra, forrásokra és vizsgálati időpontra vonatkozik. '
                 'A tényállás vagy jogi időállapot változása a következtetések felülvizsgálatát igényelheti.')
    return list(dict.fromkeys(lines))


def add_authorities(doc, data):
    if not data.legal_references:
        doc.add_paragraph('Ügyre alkalmazott bírósági döntés nincs külön rögzítve. A kutatási javaslatok nem minősülnek alkalmazott jogforrásnak.')
    for ref in data.legal_references:
        p=doc.add_paragraph();p.add_run(ref.case_number+' – '+ref.title).bold=True
        doc.add_paragraph(('Ellenőrzött és az ügyre alkalmazott hivatkozás. ' if ref.checked else 'Kutatási hivatkozás; még nincs ellenőrizve és az ügyre alkalmazva. ')+
                          'Forrás: '+(ref.source or ref.url or 'nincs megadva'))
        if ref.url:doc.add_paragraph('Elérhetőség: '+ref.url,style='Caption')
        if ref.relevance:doc.add_paragraph('Ügyre vonatkozó értékelés: '+ref.relevance)


def signoff(doc,data,metadata):
    from docx.shared import Pt
    p=doc.add_paragraph()
    p.paragraph_format.keep_together=True
    p.paragraph_format.line_spacing=1
    p.paragraph_format.space_before=Pt(6)
    p.paragraph_format.space_after=Pt(0)
    p.add_run('Kelt: '+(data.report_place+', ' if data.report_place else '')+'________________'+
              '\n\n_______________________________\n'+
              (data.report_signatory or metadata.get('approver') or 'Szakértő / aláíró neve')+
              ('\n'+data.report_issuer if data.report_issuer else ''))
