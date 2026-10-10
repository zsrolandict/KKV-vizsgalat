"""Versioned, case-specific expert opinion with portable Word/PDF layout."""
from io import BytesIO
from decimal import Decimal as D
from datetime import datetime
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .engine import LABELS, LIMITS
from .report_text import RELATIONS, hu, money, report_summary
from .report_graph import graph_image
from .opinion_support import market_paragraphs, assumptions, add_authorities, signoff

# Increment when the format changes, so old cached exports refresh.
REPORT_FORMAT_VERSION = '4'
GREEN, PALE, INK = '202020', 'F3F3F3', '202020'


def shade(cell, color):
    element = OxmlElement('w:shd'); element.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(element)


def page_field(paragraph, name):
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), name)
    paragraph._p.append(field)


def spacer(doc):
    p=doc.add_paragraph();p.add_run().font.size=Pt(1)
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=Pt(2)
    return p


def add_table(doc, headers, rows, widths, numeric=()):
    table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
    for column, width in zip(table.columns, widths): column.width = Cm(width)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = value; shade(cell, "E8E8E8")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True; run.font.color.rgb = RGBColor.from_string(INK)
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
            if index % 2 == 0: shade(cells[i], 'FAFAFA')
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '90'), ('bottom', '90'), ('left', '100'), ('right', '100')]:
                part = OxmlElement('w:'+side); part.set(qn('w:w'), value); part.set(qn('w:type'), 'dxa'); margins.append(part)
            cell._tc.get_or_add_tcPr().append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.05
                if i in numeric: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in p.runs: run.font.name = 'Garamond'; run.font.size = Pt(9.5)
    spacer(doc)
    return table


def summary_box(doc, label, text, draft=False):
    table=doc.add_table(rows=1,cols=2);table.autofit=False
    table.columns[0].width=Cm(3.2);table.columns[1].width=Cm(14.3)
    table.cell(0,0).text=label;table.cell(0,1).text=text
    for cell in table.rows[0].cells:
        shade(cell,'F5F5F5')
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(4)
            for run in p.runs:run.font.name='Garamond';run.font.size=Pt(12)
    table.cell(0,0).paragraphs[0].runs[0].bold=True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    spacer(doc)


def new_section(doc, landscape=False):
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(1.75)
    return section


def configure(doc, data, calculation, metadata, tao=False):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(1.75)
    section.header_distance = section.footer_distance = Cm(.8)
    normal = doc.styles['Normal']; normal.font.name = 'Garamond'; normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.line_spacing = 1.05; normal.paragraph_format.space_after = Pt(4)
    for name, size in [('Title', 20), ('Subtitle', 12), ('Heading 1', 13), ('Heading 2', 12), ('Heading 3', 12)]:
        style = doc.styles[name]; style.font.name = 'Garamond'; style.font.bold = name != 'Subtitle'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(GREEN); style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name == 'Heading 1' else 5); style.paragraph_format.space_after = Pt(4)
    for name in ['Header', 'Footer', 'Caption']:
        doc.styles[name].font.name = 'Garamond'; doc.styles[name].font.size = Pt(8)
        doc.styles[name].font.color.rgb = RGBColor.from_string('606060')
    # Clear theme fonts and the template's colored title border: Word and LibreOffice
    # should use the same explicit serif typography, including headings.
    for element in doc.styles.element.iter(qn('w:rFonts')):
        for attr in list(element.attrib):
            if 'theme' in attr.lower(): del element.attrib[attr]
        for script in ('ascii','hAnsi','eastAsia','cs'): element.set(qn('w:'+script),'Garamond')
    for border in list(doc.styles['Title'].element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    header = section.header.paragraphs[0]
    header.text = ('TAO' if tao else 'KKV') + ' / SZAKÉRTŐI ÁLLÁSFOGLALÁS' + ('' if metadata['approved'] else '  /  TERVEZET')
    borders = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    for key, value in [('val', 'single'), ('sz', '5'), ('color', 'C8C8C8'), ('space', '7')]: bottom.set(qn('w:'+key), value)
    borders.append(bottom); header._p.get_or_add_pPr().append(borders)
    footer = section.footer.paragraphs[0]; footer.text = f'{metadata["id"][:8]} · {metadata["version"]}. ügyverzió  |  '
    page_field(footer, 'PAGE'); footer.add_run(' / '); page_field(footer, 'NUMPAGES'); footer.add_run(' oldal')
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    settings = OxmlElement('w:updateFields'); settings.set(qn('w:val'), 'true'); doc.settings.element.append(settings)
    doc.core_properties.title = data.title+' – '+('Tao' if tao else 'KKV')+'-állásfoglalás'
    doc.core_properties.subject = f'{", ".join(map(str, sorted(getattr(data, 'years', [data.as_of.year]))))} · {metadata["version"]}. ügyverzió'
    doc.core_properties.author = data.report_issuer or data.report_signatory or metadata.get('approver') or 'Szakértői állásfoglalás'


def word_report(data, calculation, metadata):
    doc = Document(); configure(doc, data, calculation, metadata)
    summary = report_summary(data, calculation, metadata)
    root = next(c for c in data.companies if c.id == data.root)
    title=doc.add_heading('állásfoglalás',0);title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    subtitle=doc.add_paragraph('– A '+root.name+' KKV-minősítése vonatkozásában –',style='Subtitle');subtitle.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(root.name); p.runs[0].bold = True; p.runs[0].font.size = Pt(12)
    doc.add_paragraph(f'{", ".join(map(str, sorted(data.years)))}. üzleti évek • Vizsgálat napja: {data.as_of.isoformat()}', style='Caption')
    summary_box(doc, 'Jóváhagyott megállapítás' if metadata['approved'] else 'Tervezet · előzetes megállapítás', summary['conclusion'], not metadata['approved'])
    if not metadata['approved']: doc.add_paragraph('TERVEZET – szakértői jóváhagyás nélkül végleges állásfoglalásként nem használható.', style='Caption')

    doc.add_heading('1. A megbízott feladata és a vizsgálat kerete', 1); doc.add_paragraph(summary['opening'])
    add_table(doc, ['Vizsgálati adat', 'Rögzített érték'], [
        ['Megbízó', data.client or 'Nincs külön megadva'], ['Vizsgált társaság', root.name],
        ['Cégazonosító', root.registration or 'Nincs rögzítve'],
        ['Jogi profil', 'Magyar Kkv. törvény' if data.profile == 'HU' else 'EU KKV-definíció'],
        ['Háló alkalmazása', 'Évenként az üzleti év zárónapja' if data.structure_basis == 'period_end' else 'A vizsgálat napja'],
        ['Ügyverzió / szabályverzió', f'{metadata["version"]}. verzió / {calculation["rule_version"]}']], [5, 12.4])

    doc.add_heading('2. Megállapítások – gazdasági mutatók és értékhatárok', 1)
    doc.add_paragraph('A méretbesorolás az éves létszám, az éves nettó árbevétel és a mérlegfőösszeg alapján történik. A létszámfeltételnek és legalább az egyik pénzügyi feltételnek együttesen kell teljesülnie. A létszámhatár szigorú (<), a pénzügyi határérték megengedett (≤).')
    add_table(doc, ['Kategória', 'Létszám', 'Árbevétel, EUR', 'Mérleg, EUR'],
        [[LABELS[k], '< '+hu(n, 0), '≤ '+hu(t, 0), '≤ '+hu(b, 0)] for k, (n, t, b) in zip(['micro', 'small', 'medium'], LIMITS)], [5.1, 2.3, 5, 5], numeric=(1, 2, 3))
    add_table(doc, ['Év', 'Érvényesség', 'Jegyzés', 'HUF/EUR', 'Forrás / státusz'],
        [[str(y), r.date.isoformat() if r else 'Hiányzik', r.quoted.isoformat() if r and r.quoted else 'Nincs megadva',
          hu(r.value, 8) if r else '–', (r.source or 'Forrás hiányzik')+' · '+('ellenőrzött' if r.confirmed else 'ellenőrizendő') if r else 'Árfolyam hiányzik']
         for y in sorted(data.years) for r in [next((r for r in data.rates if r.year == y), None)]], [1.1, 2.5, 2.5, 2.1, 9.2], numeric=(3,))
    doc.add_paragraph('Az éves HUF-határok a rögzített árfolyamokból következnek; ellenőrizetlen árfolyamnál az értékek előzetesek.', style='Caption')
    add_table(doc, ['Év / kategória', 'Létszám', 'Árbevételhatár, ezer Ft', 'Mérleghatár, ezer Ft'],
        [[f'{y["year"]} · {LABELS[k]}', '< '+hu(n, 0), '≤ '+money(t*D(y['rate'])) if y['rate'] else 'Árfolyam hiányzik',
          '≤ '+money(b*D(y['rate'])) if y['rate'] else 'Árfolyam hiányzik']
         for y in calculation['years'] for k, (n, t, b) in zip(['micro', 'small', 'medium'], LIMITS)], [5.6, 1.6, 5.1, 5.1], numeric=(1, 2, 3))

    doc.add_heading('3. Vállalkozások közötti kapcsolatok és tulajdoni viszonyok', 1)
    doc.add_paragraph('A vizsgált társaság saját adatai 100%-kal szerepelnek. A kapcsolódó blokk teljes adatai, a releváns partner és annak kapcsolódó blokkja a megállapított aránnyal számítandó be. A partner partnerét nem adjuk automatikusan az összesítéshez; a konszolidált adatban már szereplő tételeket nem számítjuk kétszer.')
    doc.add_paragraph('Többségi vállalkozási szavazat vagy igazolt meghatározó irányítás kapcsolódást alapozhat meg. Kapcsolódás hiányában a tőke és szavazat közül a magasabb, legalább 25%-os arány partnerkapcsolatot alapozhat meg, az alkalmazható kivételekkel. A természetes személyi tulajdon és a rokonság önmagában nem dönti el a két cég minősítését.')
    for text in summary['relationships']: doc.add_paragraph(text)
    doc.add_paragraph('Az évenkénti arányokat és az igazoló indokokat a mellékletek tartalmazzák. A 0%-os tisztázandó tétel nem igazolt önállóság.', style='Caption')
    pr = data.public_review
    if any(c.kind == 'public' for c in data.companies) or pr.capital or pr.votes or pr.exception:
        doc.add_paragraph(f'Közjogi részesedés: a rögzített szakértői összesítés {hu(pr.capital)}% tőke és {hu(pr.votes)}% szavazat; {"ellenőrzött" if pr.confirmed else "ellenőrizendő"}. Indoklás: {pr.reason or "nincs rögzítve"}. '+('Kivétel alkalmazását jelölték; igazolása a rögzített indokolás szerint vizsgálandó.' if pr.exception else ''))

    doc.add_heading('Azonos / szomszédos piac és közös fellépés értékelése',2)
    for text in market_paragraphs(data): doc.add_paragraph(text)

    doc.add_heading('4. Éves összeszámítás és méretbesorolás', 1)
    annual_rows = lambda years: [[str(y['year']), hu(y['totals']['employees']), money(y['totals']['turnover']), money(y['totals']['balance']), y['label']+(' · hiányos adat' if not y['complete'] else '')] for y in years]
    add_table(doc, ['Év', 'Létszám, fő', 'Árbevétel, ezer Ft', 'Mérleg, ezer Ft', 'Éves méret'], annual_rows(calculation['years']), [1.1, 2.1, 4.2, 4.2, 5.8], numeric=(1, 2, 3))
    for y in summary['annuals']:
        doc.add_heading(f'{y["year"]}. üzleti év', 2); doc.add_paragraph(y['reason'])
    doc.add_paragraph('A pénzügyi táblázatok ezer Ft egységűek. A besorolást mindig a kerekítés előtti értékekből végezzük; a hiányzó adat nem nulla.', style='Caption')

    doc.add_heading('5. Kétéves szabály és történeti minősítés', 1)
    doc.add_paragraph('Az éves méretet és a történeti szabály szerinti hatályos kategóriát külön vizsgáljuk. A mérethatárok teljesülésének változását a rendszer két egymást követő időszak alapján követi, a dokumentált korábbi minősítés és az indokolt strukturális eseménydöntések figyelembevételével.')
    doc.add_paragraph(f'Rögzített előzmény: {LABELS[data.previous_category]}; referenciaidőpont: {data.previous_date or "hiányzik"}.' if data.previous_category else 'Külön korábbi minősítés nincs rögzítve; a történeti állapotot a vizsgált évek megfigyelései alapján vezetjük le.')
    for text in summary['history']: doc.add_paragraph(text)
    for e in data.events: doc.add_paragraph(f'{e.date}: {e.description}. Hatás és jogi indok: {e.resolution or "még nincs rendezve"}. Állapot: {"ellenőrzött" if e.confirmed else "ellenőrizendő"}; azonnali kategóriaváltás: {"jelölt" if e.immediate else "nem jelölt"}.')

    if summary['sensitivity']:
        doc.add_heading('6. Teljes összeszámítás – érzékenységvizsgálat', 1)
        doc.add_paragraph(summary['sensitivity']['explanation'])
        add_table(doc, ['Év', 'Létszám, fő', 'Árbevétel, ezer Ft', 'Mérleg, ezer Ft', 'Éves méret'], annual_rows(summary['sensitivity']['years']), [1.1, 2.1, 4.2, 4.2, 5.8], numeric=(1, 2, 3))
        for text in summary['sensitivity']['comparison']: doc.add_paragraph(text)

    if not summary['sensitivity']:
        doc.add_heading('6. Teljes összeszámítási érzékenységvizsgálat', 1)
        doc.add_paragraph('Egyetlen rögzített vállalkozás esetén a teljes összeszámítás az alapforgatókönyvvel azonos; külön összehasonlítás nem szükséges.')
    doc.add_heading('7. Összefoglaló megállapítás', 1)
    summary_box(doc, calculation['label'], summary['conclusion'], not metadata['approved'])
    doc.add_paragraph(summary['transfer_pricing'])
    if data.conclusion_notes:
        doc.add_heading('Rögzített szakértői kiegészítés', 2); doc.add_paragraph(data.conclusion_notes)

    doc.add_heading('8. Források, feltételezések és nyitott kérdések', 1)
    doc.add_heading('Feltételezések és korlátozások',2)
    for text in assumptions(data): doc.add_paragraph(text,style='List Bullet')
    sources = sorted({f.source for f in data.financials if f.source} | {o.source for o in data.ownerships if o.source} | {d.source for d in data.decisions if d.source} | {f.source for f in data.families if f.source})
    if sources:
        doc.add_heading('Rögzített igazoló források', 2)
        for source in sources: doc.add_paragraph(source, style='List Bullet')
    else: doc.add_paragraph('Igazoló forrás nincs rögzítve.')
    if calculation['blockers']:
        doc.add_heading('Véglegesítés előtt rendezendő kérdések', 2)
        groups = {}
        for b in calculation['blockers']:
            message = b['message']
            if len(message) > 5 and message[:4].isdigit() and message[4:6] == ': ': message = message[6:]
            key = b['code'], b.get('target'), message
            if key not in groups: groups[key] = [b, [], message]
            if b['year']: groups[key][1].append(str(b['year']))
        for b, years, message in groups.values():
            doc.add_paragraph((', '.join(years)+': ' if years else '')+message, style='List Bullet')
            if b.get('action'): doc.add_paragraph('Teendő: '+b['action'], style='Caption')
    else: doc.add_paragraph('Nincs nyitott számítási vagy adatellenőrzési kérdés. Ez nem helyettesíti a szakértői jóváhagyást.')
    for warning in calculation['warnings']: doc.add_paragraph(warning)

    doc.add_heading('9. Jogi keret és szakértői jóváhagyás', 1)
    doc.add_paragraph(f'Ellenőrzött jogi időállapot: {data.law_date or "nincs rögzítve"}. Forrás: {data.law_source or "nincs rögzítve"}. Történeti alkalmazhatóság: {data.law_applicability or "nem igényel külön indokot"}.')
    doc.add_paragraph('Magyar profil: a 2004. évi XXXIV. törvény a kis- és középvállalkozásokról, fejlődésük támogatásáról. EU-profil: a 2003/361/EK ajánlás és a vizsgálat céljára alkalmazandó uniós rendelkezések. Az alkalmazandó időállapotot, kivételeket és a szabályok konkrét ügyre való alkalmazhatóságát a jóváhagyó szakértő ellenőrzi.')
    doc.add_paragraph('Transzferár esetén az adott adóév társaságiadó- és nyilvántartási szabályai, a kapcsolt ügyletek és a mentességek külön vizsgálandók. A méretkategória önmagában nem állapít meg ügyletszintű nyilvántartási kötelezettséget.')
    doc.add_heading('Alkalmazott jogszabályok és bírósági döntések',2)
    add_authorities(doc,data)
    if metadata['approved']:
        stamp = datetime.fromisoformat(metadata['approved_at']).astimezone(ZoneInfo('Europe/Budapest')).strftime('%Y.%m.%d. %H:%M')
        doc.add_paragraph('Jóváhagyó szakértő: '+metadata['approver'])
        add_table(doc, ['Jóváhagyási adat', 'Rögzített érték'], [['Szakértő', metadata['approver']], ['Időpont', stamp+' (Europe/Budapest)'], ['Jóváhagyott ügyverzió', str(metadata['version'])]], [5, 12.4])
        doc.add_paragraph('A rendszerbeli jóváhagyás nem elektronikus aláírás. A dokumentum aláírása a szolgáltató munkafolyamata szerint történik.', style='Caption')
    else: doc.add_paragraph('Szakértői jóváhagyás még nem történt. A dokumentum tervezet.')

    signoff(doc,data,metadata)
    new_section(doc, landscape=True)
    doc.add_heading('1. melléklet – részletes éves gazdasági mutatók', 1)
    doc.add_paragraph('A táblázat a jogcím szerinti súlyozás után figyelembe vett értékeket tartalmazza, nem a súlyozás előtti egyedi beszámolót.', style='Caption')
    for year in calculation['years']:
        doc.add_heading(f'{year["year"]}. üzleti év · zárónap: {year["closing"]}', 2)
        rows = [[r['name'], RELATIONS[r['relation']], hu(r['percent'])+'%', hu(r['employees']), money(r['turnover']), money(r['balance'])] for r in year['rows']]
        rows.append(['Összesen', 'Teljes adatkör' if year['complete'] else 'Hiányos adatkör', '', hu(year['totals']['employees']), money(year['totals']['turnover']), money(year['totals']['balance'])])
        table = add_table(doc, ['Vállalkozás', 'Jogcím', 'Arány', 'Létszám, fő', 'Árbevétel, ezer Ft', 'Mérleg, ezer Ft'], rows, [7.1, 4.7, 1.8, 2.6, 5, 4.9], numeric=(2, 3, 4, 5))
        for cell in table.rows[-1].cells:
            shade(cell, PALE)
            for p in cell.paragraphs:
                for r in p.runs: r.bold = True
        doc.add_paragraph('„–”: nincs külön beszámított összeg; a hiányzó, kizárt vagy konszolidációban már szereplő tételt a jogcím és az indokolás alapján kell értelmezni.', style='Caption')

    new_section(doc)
    doc.add_heading('2. melléklet – kapcsolati indokok és forrásadatok', 1)
    for year in calculation['years']:
        doc.add_heading(f'{year["year"]}. üzleti év · beszámítás indokai', 2)
        for row in year['rows']:
            p = doc.add_paragraph(); p.add_run(row['name']+' – ').bold = True
            p.add_run(RELATIONS[row['relation']]+f' ({hu(row["percent"])}%). '+row['reason'])
    names = {c.id: c.name for c in [*data.companies, *data.persons]}
    if data.decisions:
        doc.add_heading('Szakértői kapcsolati döntések', 2)
        for d in data.decisions:
            p = doc.add_paragraph(); p.add_run(f'{names[d.first]} ↔ {names[d.second]}: ').bold = True
            p.add_run(f'{RELATIONS[d.relation]}; {d.reason or "indoklás hiányzik"}. Forrás: {d.source or "hiányzik"}. Érvényesség: {d.start or "nincs kezdő korlát"} – {d.end or "nincs záró korlát"}; {"ellenőrzött" if d.confirmed else "ellenőrizendő"}.')
            if d.basis == 'persons': doc.add_paragraph(f'Közös fellépés: {d.acting_together or "tisztázandó"}. Piaci viszony: {d.market or "tisztázandó"}.', style='Caption')
    if data.families:
        doc.add_heading('Rokonsági tényállás', 2)
        for f in data.families: doc.add_paragraph(f'{names[f.first]} – {names[f.second]}: {f.relationship}; forrás: {f.source or "nincs megadva"}.')
        doc.add_paragraph('A rokonsági adat önmagában nem eredményez automatikus összeszámítást.', style='Caption')
    doc.add_heading('Beszámolóadatok eredete', 2)
    add_table(doc, ['Cég / év', 'Időszak / elfogadás', 'Módszer / forrás'],
        [[f'{names[f.company]} · {f.year}', f'{f.start or "?"} – {f.end or "?"}\nElfogadás: {f.accepted or "nincs"}',
          f'{f.currency} · {f.employment_method}\n{f.source or "Forrás hiányzik"}'+('\nÉvesített becslés' if f.estimated else '')]
         for f in sorted(data.financials, key=lambda f: (f.year, names[f.company]))], [5.5, 5.2, 6.7])

    new_section(doc, landscape=True)
    for index, year in enumerate(calculation['years']):
        if index: doc.add_page_break()
        doc.add_heading(f'3. melléklet – cégháló · {year["year"]}', 1)
        doc.add_paragraph('Az ügyverzióban mentett elrendezéssel. A minősítés és a tulajdoni/döntési élek az adott hálóidőponthoz tartoznak.', style='Caption')
        png = graph_image(data, calculation, year['year'])
        from PIL import Image
        with Image.open(png) as image: ratio = image.width/image.height
        png.seek(0)
        paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture = paragraph.add_run().add_picture(png, width=Cm(min(25, 12.5*ratio)))
        picture._inline.docPr.set('descr', f'{year["year"]}. évi cégháló, a mentett ügyverzió pozícióival.')
        doc.add_paragraph('A kapcsolati minősítés részletes indokolását és a teljes cégneveket a 2. melléklet tartalmazza. A rokonsági él nem önálló kapcsolódási szabály.', style='Caption')
    output = BytesIO(); doc.save(output)
    return output.getvalue()
