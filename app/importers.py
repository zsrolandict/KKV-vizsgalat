from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile
import openpyxl
from .models import Assessment


def safe_workbook(raw):
    with ZipFile(BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist()) > 80*1024*1024:
            raise ValueError('A kicsomagolt Excel túl nagy.')
    return openpyxl.load_workbook(BytesIO(raw), data_only=True, read_only=True)


def import_excel(raw, filename):
    wb = safe_workbook(raw)
    try:
        if 'Tulajdonosi struktúra' in wb.sheetnames:
            return import_original(wb, filename)
        if 'Vállalkozások' in wb.sheetnames and 'Pénzügyi adatok' in wb.sheetnames:
            return import_template(wb, filename)
        raise ValueError('Nem támogatott Excel-szerkezet. Használja az adatbekérő sablont vagy az eredeti KKV-munkafüzetet.')
    finally:
        wb.close()


def iso(value):
    if isinstance(value,(date,datetime)):
        return value.date().isoformat() if isinstance(value,datetime) else value.isoformat()
    return str(value).strip() if value else None


def import_original(wb, filename):
    s=wb['Tulajdonosi struktúra']
    if s.max_row>5000 or s.max_column>200:
        raise ValueError('Az Excel mérete meghaladja a támogatott importhatárt.')
    company_rows=[]
    for row in s.iter_rows(min_row=4,values_only=True):
        if len(row)>4 and isinstance(row[0],(int,float)) and row[0]>=1 and isinstance(row[1],str) and isinstance(row[2],(int,float)):
            company_rows.append(row)
    if not company_rows:
        raise ValueError('Nem található vállalkozási alapadat.')
    companies=[{'id':f'c{i+1}','name':r[1]} for i,r in enumerate(company_rows)]
    cids={c['name']:c['id'] for c in companies}
    financials=[];years=[]
    for col in [2,5,8,11]:
        y=s.cell(2,col+1).value
        if not isinstance(y,int):continue
        for r,c in zip(company_rows,companies):
            if len(r)>col+2 and all(isinstance(r[col+k],(int,float)) for k in range(3)):
                financials.append({'company':c['id'],'year':y,'employees':str(r[col]),
                    'turnover':str(Decimal(str(r[col+1]))*1000),'balance':str(Decimal(str(r[col+2]))*1000),
                    'start':f'{y}-01-01','end':f'{y}-12-31','source':f'Excel-import: {filename}'})
    rates=[];decisions=[]
    sheets=[wb[n] for n in wb.sheetnames if n.startswith('Pénzügyi adatok_')]
    for ps in sheets:
        y=ps['H2'].value
        if not isinstance(y,int):continue
        years.append(y)
        r=ps['K17'].value;day=ps['K16'].value
        if isinstance(r,(int,float)) and day:
            rates.append({'year':y,'date':iso(day),'value':str(r),'source':f'Excel-import: {filename}; MNB-ellenőrzés szükséges.'})
    current=max(sheets,key=lambda ps:ps['H2'].value or 0) if sheets else None
    if current:
        for row in current.iter_rows(min_row=5,max_row=min(current.max_row,1000),values_only=True):
            name=row[1] if len(row)>1 else None;pct=row[2] if len(row)>2 else None
            if name not in cids or not isinstance(pct,(int,float)):continue
            rel='linked' if pct>0.5 else 'partner' if pct>=0.25 else 'independent'
            decisions.append({'id':str(uuid4()),'first':companies[0]['id'],'second':cids[name],
                'relation':rel,'percent':str(Decimal(str(pct))*100) if rel=='partner' else 0,
                'reason':'Importált kézi befolyás; szakértői kapcsolati döntésként ellenőrizendő.',
                'source':f'{filename}: befolyás oszlop','confirmed':False})
    persons={};ownerships=[];current_id=None
    for row in s.iter_rows(min_row=4,values_only=True):
        if len(row)>2 and row[1] in cids and isinstance(row[2],(int,float)):
            current_id=cids[row[1]];continue
        if current_id and isinstance(row[0],(int,float)) and 0<=row[0]<=1 and isinstance(row[1],str):
            name=row[1];pid=persons.setdefault(name,f'p{len(persons)+1}')
            ownerships.append({'id':str(uuid4()),'owner':pid,'company':current_id,
                'capital':str(Decimal(str(row[0]))*100),'votes':str(Decimal(str(row[0]))*100),
                'source':f'{filename}: tulajdonosi struktúra; a szavazati azonosság ellenőrizendő.'})
    return Assessment.model_validate({'title':'Importált KKV-vizsgálat','client':'',
        'purpose':'KKV-minősítés','as_of':date.today().isoformat(),'root':companies[0]['id'],
        'years':sorted(set(years)) or sorted({f['year'] for f in financials})[-2:],
        'companies':companies,'persons':[{'id':v,'name':k} for k,v in persons.items()],
        'ownerships':ownerships,'decisions':decisions,'financials':financials,'rates':rates,
        'assumptions':'Excel-import. Az adatok, elfogadási időpontok, szavazati jogok és kapcsolati besorolások ellenőrzendők. Az eredeti pénzügyi egység ezer Ft, itt HUF.'})


def import_template(wb, filename):
    def rows(name):
        if name not in wb.sheetnames:return []
        if wb[name].max_row>5000:raise ValueError('Túl sok sor az Excelben.')
        return [r for r in wb[name].iter_rows(min_row=2,values_only=True) if r[0] is not None]
    companies=[{'id':str(r[0]),'name':str(r[1]),'registration':str(r[2] or ''),
                'country':str(r[3] or 'HU'),'activity':str(r[4] or '')} for r in rows('Vállalkozások')]
    persons=[{'id':str(r[0]),'name':str(r[1])} for r in rows('Személyek')]
    finances=[{'company':str(r[0]),'year':int(r[1]),'employees':r[2],'turnover':r[3],
               'balance':r[4],'accepted':iso(r[5]),'end':iso(r[6]),'source':str(r[7] or f'Excel-import: {filename}')} for r in rows('Pénzügyi adatok')]
    ownerships=[{'id':str(uuid4()),'owner':str(r[0]),'company':str(r[1]),'capital':r[2],
                'votes':r[3],'source':str(r[4] or filename)} for r in rows('Kapcsolatok')]
    rates=[{'year':int(r[0]),'date':iso(r[1]),'quoted':iso(r[2]),'value':r[3],
            'source':str(r[4] or filename),'confirmed':False} for r in rows('Árfolyamok')]
    if not companies or not finances:raise ValueError('Vállalkozás és pénzügyi adat szükséges.')
    return Assessment.model_validate({'title':'Importált KKV-vizsgálat','as_of':date.today().isoformat(),
        'root':companies[0]['id'],'years':sorted({f['year'] for f in finances}),
        'companies':companies,'persons':persons,'ownerships':ownerships,'financials':finances,'rates':rates})


def template_bytes():
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name,headers in [
      ('Vállalkozások',['Azonosító','Név','Cégjegyzékszám','Ország','Tevékenység']),
      ('Személyek',['Azonosító','Név']),
      ('Pénzügyi adatok',['Cégazonosító','Év','Létszám','Árbevétel (HUF)','Mérlegfőösszeg (HUF)','Elfogadás','Zárás','Forrás']),
      ('Kapcsolatok',['Tulajdonosazonosító','Cégazonosító','Tőke (%)','Szavazat (%)','Forrás']),
      ('Árfolyamok',['Év','Érvényesség','Jegyzés','HUF/EUR','Forrás'])]:
        s=wb.create_sheet(name);s.append(headers);s.freeze_panes='A2'
        for cell in s[1]:cell.font=openpyxl.styles.Font(bold=True,color='FFFFFF');cell.fill=openpyxl.styles.PatternFill('solid',fgColor='234B3F')
        for col in range(1,len(headers)+1):s.column_dimensions[openpyxl.utils.get_column_letter(col)].width=25
    out=BytesIO();wb.save(out);return out.getvalue()
