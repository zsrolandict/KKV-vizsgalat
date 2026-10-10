"""Exercise Tao corrections, immediate guidance, snapshots and persisted network."""
from playwright.sync_api import expect
from docx import Document


def validate_tao_workflow(page,out):
    request=page.context.request
    session=request.get('http://127.0.0.1:8011/api/auth/me').json()
    h={'X-CSRF-Token':session['csrf']}
    d={'title':'Tao vezetett ellenőrzés','as_of':'2025-12-31','scope':'A felsorolt cégek kapcsoltsága',
       'law_date':'2025-12-31','law_source':'Fiktív ellenőrzött tesztforrás',
       'companies':[{'id':'a','name':'Alfa Kft.'},{'id':'b','name':'Béta Kft.'}],
       'persons':[{'id':'p1','name':'Első tulajdonos'},{'id':'p2','name':'Második tulajdonos'}],
       'voting_facts':[{'id':'f1','owner':'p1','company':'a','capital':'100','reviewed_as_of':'2025-12-31','evidence':{'source':'Okirat'}},
                       {'id':'f2','owner':'p2','company':'b','capital':'100','reviewed_as_of':'2025-12-31','evidence':{'source':'Okirat'}}]}
    r=request.post('http://127.0.0.1:8011/api/tao/cases',data={'data':d},headers=h)
    assert r.status==200,r.text();cid=r.json()['id']
    page.goto(f'http://127.0.0.1:8011/tao?case={cid}')
    page.locator('.tabs [data-tab=review]').click()
    expect(page.locator('#pending-count')).to_have_text('1')
    expect(page.locator('.review-card')).to_contain_text('Mit kell megadnod?')
    page.locator('[data-action=resolve-pair]').click()
    page.locator('#editor [name=result]').select_option('not_related')
    expect(page.locator('#editor-recalculation')).to_contain_text('Újraszámítás szükséges')
    page.locator('#editor [name=basis]').fill('Ellenőrzött Tao. 4. § 23. pont')
    page.locator('#editor [name=reason]').fill('A vállalt kör minden releváns jogalapját ellenőriztem; nincs kapcsoltság.')
    page.locator('#editor [name=source]').fill('Ellenőrzött okiratok')
    page.locator('#editor [name=relevant_grounds_reviewed]').check()
    # Failure keeps the draft form so a retry is possible.
    page.route('**/api/tao/cases/'+cid,lambda route:route.fulfill(status=409,json={'detail':'Verzióütközés – teszt'}) if route.request.method=='PUT' else route.continue_())
    page.locator('#editor [data-save-mode=reviewed]').click()
    expect(page.locator('#editor-error')).to_contain_text('Verzióütközés')
    assert page.locator('#editor [name=reason]').input_value().startswith('A vállalt')
    page.unroute('**/api/tao/cases/'+cid)
    page.locator('#editor [data-save-mode=reviewed]').click();expect(page.locator('#editor')).not_to_be_visible()
    expect(page.locator('#pending-count')).to_have_text('0')
    expect(page.locator('.tabs [data-tab=review]')).to_have_text('Ellenőrzés (0)')
    expect(page.locator('#save-feedback')).to_contain_text('0 ellenőrizendő')
    page.reload();page.locator('.tabs [data-tab=review]').click();expect(page.locator('#pending-count')).to_have_text('0')
    page.locator('.tabs [data-tab=data]').click();page.locator('[data-action=family]').click()
    page.locator('#editor [name=relationship]').select_option('sibling')
    expect(page.locator('#editor-impact')).to_contain_text('összeszámítja')
    page.locator('#editor [name=source]').fill('Rokonsági nyilatkozat');page.locator('#editor [name=confirmed]').check()
    page.locator('#editor [type=submit]').click();expect(page.locator('#editor')).not_to_be_visible()
    # New evidence invalidates an unchanged prior confirmation, and a family signal appears.
    page.locator('.tabs [data-tab=review]').click();expect(page.locator('#pending-count')).to_have_text('1')
    expect(page.locator('#save-feedback')).to_contain_text('Változott')
    page.locator('.tabs [data-tab=graph]').click()
    node=page.locator('[data-network-node=a]');before=node.get_attribute('transform')
    edge=page.locator('#tao-network path[marker-end]').first.get_attribute('d')
    box=node.bounding_box();page.mouse.move(box['x']+40,box['y']+20);page.mouse.down();page.mouse.move(box['x']+130,box['y']+100,steps=10);page.mouse.up()
    assert page.locator('[data-network-node=a]').get_attribute('transform')!=before
    assert page.locator('#tao-network path[marker-end]').first.get_attribute('d')!=edge
    position=page.locator('[data-network-node=a]').get_attribute('transform')
    page.locator('[data-action=save-graph]').click();expect(page.locator('[data-action=save-graph]')).to_be_disabled()
    page.reload();page.locator('.tabs [data-tab=graph]').click();expect(page.locator('[data-network-node=a]')).to_have_attribute('transform',position)
    node=page.locator('[data-network-node=a]');node.focus();node.press('ArrowRight');assert node.get_attribute('transform')!=position
    page.locator('[data-action=save-graph]').click();expect(page.locator('[data-action=save-graph]')).to_be_disabled()
    for kind in ['svg','png']:
        with page.expect_download() as event:page.locator(f'[data-action=graph-download][data-kind={kind}]').click()
        path=out/f'tao-ceghalo.{kind}';event.value.save_as(path)
        assert path.stat().st_size>500
        if kind=='png':assert path.read_bytes().startswith(b'\x89PNG')
    for width in [360,768,1440]:
        page.set_viewport_size({'width':width,'height':1000});assert not page.evaluate('document.documentElement.scrollWidth>innerWidth'),f'Tao overflow {width}'
    page.screenshot(path=str(out/'tao-uj-ceghalo.png'),full_page=True)
    page.set_viewport_size({'width':1440,'height':1050})
    with page.expect_download() as event:page.locator('[data-action=export][data-kind=docx]').click()
    path=out/'tao-allasfoglalas.docx';event.value.save_as(path);doc=Document(path)
    assert len(doc.inline_shapes)==1
    assert 'Vezetői megállapítások' in '\n'.join(p.text for p in doc.paragraphs)
    with page.expect_download(timeout=90000) as event:page.locator('[data-action=export][data-kind=pdf]').click()
    path=out/'tao-allasfoglalas.pdf';event.value.save_as(path);assert path.read_bytes().startswith(b'%PDF')


def validate_tao_demo_decisions(page,out):
    page.goto('http://127.0.0.1:8011/tao');page.locator('[data-action=demo]').click()
    page.locator('.tabs [data-tab=review]').click();expect(page.locator('#pending-count')).to_have_text('16')
    # A plain draft can be saved, with an explicit explanation of why the counter remains.
    page.locator('[data-action=resolve-pair]').first.click()
    page.locator('#editor [name=legal_ground]').select_option('abc')
    expect(page.locator('[data-select-help=legal_ground]')).to_contain_text('telephelyi')
    page.locator('#editor [value=draft]').click();expect(page.locator('#editor')).not_to_be_visible()
    expect(page.locator('#pending-count')).to_have_text('16')
    expect(page.locator('#save-feedback')).to_contain_text('tervezetként')
    # The reviewed action closes a pair without a hidden extra checkbox being required.
    page.locator('[data-action=resolve-pair]').first.click()
    page.locator('#editor [name=result]').select_option('not_related')
    page.locator('#editor [name=basis]').fill('Tao. releváns rendelkezései, mintateszt')
    page.locator('#editor [name=reason]').fill('A vállalt kör minden releváns jogalapja ellenőrizve; nem kapcsolt.')
    page.locator('#editor [name=source]').fill('Fiktív ellenőrzött mintairat')
    page.locator('#editor [name=relevant_grounds_reviewed]').check()
    page.locator('#editor [data-save-mode=reviewed]').click();expect(page.locator('#editor')).not_to_be_visible()
    expect(page.locator('#pending-count')).to_have_text('15');expect(page.locator('.tabs [data-tab=review]')).to_have_text('Ellenőrzés (15)')
    page.reload();page.locator('.tabs [data-tab=review]').click();expect(page.locator('#pending-count')).to_have_text('15')
    # An explicitly uncertain reviewed finding is saved without falsely closing the question.
    page.locator('[data-action=resolve-pair]').first.click()
    page.locator('#editor [name=result]').select_option('undetermined')
    page.locator('#editor [name=reason]').fill('Hiányzik az irányítási szerződés.')
    page.locator('#editor [name=missing]').fill('Kérd be az irányítási szerződést a társaságtól.')
    page.locator('#editor [data-save-mode=reviewed]').click();expect(page.locator('#editor')).not_to_be_visible()
    expect(page.locator('#pending-count')).to_have_text('15');expect(page.locator('#save-feedback')).to_contain_text('további tisztázást')
    # Legal research candidates are editable; they are not silently applied.
    page.locator('details').filter(has_text='Kutatási kiindulópontok').locator('summary').click()
    page.locator('[data-action=legal-reference][data-candidate="0"]').click()
    expect(page.locator('#editor [name=case_number]')).to_have_value('C-110/13')
    page.locator('#editor [name=source]').fill('A hivatalos döntés vizsgált bekezdései – teszt')
    page.locator('#editor [name=relevance]').fill('Ügyre vonatkozó tényállási összevetés – teszt')
    page.locator('#editor [name=checked]').check();page.locator('#editor [type=submit]').click()
    expect(page.locator('#editor')).not_to_be_visible();expect(page.get_by_text('Ellenőrzött, ügyre alkalmazott',exact=True)).to_have_count(1)
    page.get_by_role('button',name='Ügyadatok',exact=True).click();page.locator('#editor [name=market_analysis]').fill('A szolgáltatások és a földrajzi piac ügyre szabott értékelése – teszt.')
    page.locator('#editor [name=assumptions]').fill('A megadott adatok teljességét a megbízói nyilatkozat alapján vizsgáljuk – teszt.')
    page.locator('#editor [name=report_issuer]').fill('Minta szakértői iroda');page.locator('#editor [type=submit]').click();expect(page.locator('#editor')).not_to_be_visible()
    page.locator('[data-action=pdf-tools]').click();expect(page.locator('#editor')).to_contain_text('LibreOffice');expect(page.locator('#editor')).to_contain_text('Poppler');page.locator('#editor [type=submit]').click()
    page.locator('[data-action=resolve-pair]').first.click();page.locator('#editor [name=legal_ground]').select_option('management')
    expect(page.locator('[data-select-help=legal_ground]')).to_contain_text('közös ügyvezető')
    page.screenshot(path=str(out/'tao-jogalap-magyarazat.png'),full_page=True)
    page.on('dialog',lambda d:d.accept());page.locator('#editor [data-action=close]').first.click()
    with page.expect_download() as event:page.locator('[data-action=export][data-kind=docx]').click()
    path=out/'tao-reszletes-minta.docx';event.value.save_as(path)
    doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)
    assert 'C-110/13' in text and 'földrajzi piac' in text and 'Minta szakértői iroda' in text
