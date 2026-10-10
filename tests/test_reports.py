from io import BytesIO
import unittest
from docx import Document
from PIL import Image
from app.engine import calculate
from app.report_text import report_summary, annual_reason
from app.report_graph import graph_positions, graph_image
from app.reports import word_report
from .helpers import sample, extended


class ReportTests(unittest.TestCase):
    def test_reason_uses_exact_limits_and_financial_or(self):
        data=sample().model_dump(mode='json')
        for f in data['financials']:
            f.update(employees='9.9',turnover='800000001',balance='800000000')
        from app.models import Assessment
        model=Assessment.model_validate(data);year=calculate(model)['years'][0]
        self.assertEqual(year['category'],'micro')
        reason=annual_reason(year)
        self.assertIn('9,9',reason);self.assertIn('meghaladja',reason)
        self.assertIn('legalább az egyik pénzügyi feltétel teljesül',reason)
        for f in data['financials']:f.update(employees='10',balance='800000001')
        year=calculate(Assessment.model_validate(data))['years'][0]
        self.assertEqual(year['category'],'small')
        self.assertIn('a létszámhatár nem teljesül',annual_reason(year))

    def test_unresolved_opinion_is_qualified_and_sensitivity_is_separate(self):
        model=extended(['b'],[]);calc=calculate(model)
        summary=report_summary(model,calc,{'approved':False,'version':1})
        self.assertIn('még nem igazolható',summary['conclusion'])
        self.assertIn('nem igazolt kizárás',' '.join(summary['relationships']))
        self.assertIn('nem önálló jogi minősítés',summary['sensitivity']['explanation'])
        doc=Document(BytesIO(word_report(model,calc,{'id':'test','approved':False,'version':1})))
        self.assertIn('TERVEZET',doc.sections[0].header.paragraphs[0].text)
        self.assertEqual(len(doc.inline_shapes),2)
        self.assertTrue(any('részletes éves' in p.text for p in doc.paragraphs))

    def test_graph_uses_saved_positions_and_is_bounded(self):
        data=extended(['b'],[('a','b',60)])
        from app.models import GraphPosition
        data.graph_positions={'a':GraphPosition(x=320,y=250),'b':GraphPosition(x=5000,y=5000)}
        self.assertEqual(graph_positions(data)['a'],(320,250))
        with Image.open(graph_image(data,calculate(data),2025)) as image:
            self.assertEqual(image.format,'PNG')
            self.assertLessEqual(max(image.size),2800)
            self.assertLessEqual(image.width*image.height,12000000)

    def test_approved_text_matches_preview_and_has_approval_identity(self):
        data=sample();calc=calculate(data)
        meta={'id':'test','approved':True,'version':3,'approved_at':'2026-10-10T12:00:00+00:00','approver':'Minta Szakértő'}
        summary=report_summary(data,calc,meta)
        doc=Document(BytesIO(word_report(data,calc,meta)))
        table_text=' '.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn(summary['conclusion'],table_text)
        self.assertIn('Minta Szakértő',' '.join(p.text for p in doc.paragraphs))
        self.assertNotIn('TERVEZET',doc.sections[0].header.paragraphs[0].text)
        self.assertIn('9. Jogi keret',' '.join(p.text for p in doc.paragraphs))
