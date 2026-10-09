import copy
import os
import secrets
import shutil
from docx import Document
import tempfile
import unittest
from datetime import date
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import db
from app.main import app
from app.pdf_import import parse_text
from app.pdf_api import Review, draft
from app.tao_models import TaoAssessment, VotingFact
from app.tao_engine import calculate, active

TEXT = '''Minta Holding Zrt.
Cégjegyzékszám: 01 10 123456
Cégadatlap 2026. 10. 09-i hatállyal
1(10). A részvényesek adatai
1(10)/1 [01 09 654321 ] Minta Vagyonkezelő Korlátolt Felelősségű Társaság
 A cég tagja e jogállását bizalmi vagyonkezelőként szerezte meg.
 A részvényes egyedüli részvényes.
 Bejegyzés kelte: 2024.01.23.
 Hatályos: 2022.01.01. - 2023.09.15.
3(10). A részvényátruházás korlátozásának jelzése
13. A cégjegyzésre jogosult(ak) adatai
13/1 Minta Anna (an: Minta Éva) ügyvezető (vezető tisztségviselő)
 Hatályos: 2024.01.01. - ...
20. A cég statisztikai számjele
97. Pénzügyi modul
        2025. év         2024. év
Beszámolási időszak 2025.01.01 -   2024.01.01 -
Értékek: Ezer HUF-ban 2025.12.31   2024.12.31
Értékesítés nettó árbevétele   12 345       23 456
Eszközök összesen   100 123       90 234
Létszám: 4 fő
'''

def tao():
    return TaoAssessment(title='Minta', as_of='2024-12-31', companies=[dict(id='a', name='A'), dict(id='b', name='B')])


class ParserTests(unittest.TestCase):
    def test_source_dates_and_missing_annual_staff(self):
        d=parse_text(TEXT)
        self.assertEqual(d['registration'], '01-10-123456')
        self.assertEqual(len(d['owners']),1)
        self.assertEqual(d['owners'][0]['valid_to'],'2023-09-15')
        self.assertEqual(d['owners'][0]['registered_on'],'2024-01-23')
        self.assertTrue(d['owners'][0]['trustee'])
        self.assertEqual(d['reference_employees']['value'],'4')
        self.assertIsNone(d['financials'][0]['employees'])
        self.assertEqual(d['financials'][0]['turnover'],'12345')
        self.assertEqual(len(d['leaders']),1)

    def test_qualitative_vote_is_not_invented_and_collection_not_ownership(self):
        d=parse_text(TEXT.replace('A részvényes egyedüli részvényes.','Szavazati jog mértéke meghaladja az 50%-ot.\n Befolyás mértéke: 30% (Opten gyűjtés)'))
        o=d['owners'][0]
        self.assertEqual(o['vote_bound'],'over_half'); self.assertIsNone(o['votes']); self.assertIsNone(o['capital'])
        self.assertEqual(o['raw_influence'],'30')

    def test_duplicate_company_source_preserves_reviewed_name_and_ownership(self):
        first=parse_text(TEXT);first['name']='Kézzel ellenőrzött név';first['document_id']='a'
        second=parse_text(TEXT);second['document_id']='b'
        p=Review(module='kkv',title='Minta',as_of='2024-12-31',root=first['registration'],years=[2024],documents=[first,second],confirmed=True)
        data,_=draft(p)
        self.assertEqual(next(c.name for c in data.companies if c.id==data.root),'Kézzel ellenőrzött név')
        self.assertEqual(len(data.ownerships),1)

    def test_copied_reference_headcount_requires_method_review(self):
        from app.engine import calculate as kkv_calculate
        doc=parse_text(TEXT);doc['financials'][1]['employees']='4'
        p=Review(module='kkv',title='Minta',as_of='2024-12-31',root=doc['registration'],years=[2024],documents=[doc],confirmed=True)
        data,_=draft(p)
        self.assertEqual(data.financials[0].employment_method,'estimate')
        self.assertTrue(any(b['code']=='employment_method' for b in kkv_calculate(data)['blockers']))

    def test_no_text_and_other_pdf_rejected(self):
        for text in [' ', 'Different document '*20]:
            with self.assertRaises(ValueError): parse_text(text)

    def test_build_separate_drafts_and_exact_huf_conversion(self):
        d=parse_text(TEXT); d['document_id']='source'
        p=Review(module='kkv',title='Minta',as_of='2026-10-09',root=d['registration'],years=[2024,2025],documents=[d],confirmed=True)
        kkv,_=draft(p)
        self.assertEqual(str(kkv.financials[0].turnover),'12345000')
        self.assertIsNone(kkv.financials[0].employees)
        self.assertFalse(kkv.decisions[0].confirmed)
        self.assertEqual(kkv.financials[0].end,date(2025,12,31))
        self.assertEqual(kkv.ownerships[0].end,date(2023,9,14))
        p.module='tao'; t,_=draft(p)
        self.assertIsInstance(t,TaoAssessment)
        self.assertFalse(calculate(t)['complete'])


class TaoTests(unittest.TestCase):
    def fact(self,**kw):
        return VotingFact(id='f',owner='a',company='b',capital=60,valid_from='2024-01-01',**kw)

    def test_signal_never_automatic_legal_result(self):
        d=tao(); d.voting_facts=[self.fact()];r=calculate(d)['rows'][0]
        self.assertEqual(r['result'],'undetermined');self.assertTrue(r['signals'])
        d.voting_facts=[VotingFact(id='f',owner='a',company='b',capital=50,valid_from='2024-01-01')]
        self.assertFalse(calculate(d)['rows'][0]['signals'])

    def test_expert_priority_and_conflicts(self):
        d=tao();d.voting_facts=[self.fact(),VotingFact(id='x',owner='a',company='b',vote_mode='expert',votes=40,reason='Eltérő jogosultság',valid_from='2024-01-01')]
        self.assertFalse(calculate(d)['rows'][0]['signals'])
        d.voting_facts.append(VotingFact(id='y',owner='a',company='b',vote_mode='expert',votes=60,reason='Másik forrás',valid_from='2024-01-01'))
        self.assertTrue(calculate(d)['rows'][0]['missing'])

    def test_legal_end_and_unknown_start(self):
        f=VotingFact(id='f',owner='a',company='b',capital=60,valid_from='2022-01-01',valid_to='2023-09-15',deletion_registered_on='2024-01-23')
        self.assertFalse(active(f,date(2024,1,1)))
        f=VotingFact(id='u',owner='a',company='b',capital=None)
        self.assertFalse(active(f,date(2024,1,1)))
        self.assertTrue(calculate(tao())['counts']['undetermined'])

    def test_negative_requires_source_and_grounds(self):
        d=tao().model_dump(mode='json');d['decisions']=[dict(first='a',second='b',as_of='2024-12-31',result='not_related',confirmed=True,basis='Tao',reason='Vizsgálat',evidence={'source':'Nyilatkozat'})]
        with self.assertRaises(ValueError):TaoAssessment.model_validate(d)
        d['decisions'][0]['relevant_grounds_reviewed']=True
        self.assertEqual(calculate(TaoAssessment.model_validate(d))['rows'][0]['result'],'not_related')


class PDFAPITests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.env=patch.dict(os.environ,{'KKV_DATA_DIR':self.tmp.name});self.env.start();db.initialize()
        self.client=TestClient(app);self.password=secrets.token_urlsafe(24)
        r=self.client.post('/api/auth/setup',json=dict(name='Teszt',username='tester',password=self.password))
        self.h={'X-CSRF-Token':r.json()['csrf']}
    def tearDown(self):self.client.close();self.env.stop();self.tmp.cleanup()
    def upload(self):
        with patch('app.pdf_api.extract',return_value=parse_text(TEXT)):
            r=self.client.post('/api/pdf/batches',files=[('files',('sample.pdf',b'%PDF-1.4\nsource','application/pdf'))],headers=self.h)
        self.assertEqual(r.status_code,200,r.text);return r.json()
    def payload(self,b,module='tao'):
        d=b['files'][0]['candidate'];d['document_id']=b['files'][0]['id']
        return dict(module=module,title='Minta',as_of='2024-12-31',documents=[d],confirmed=True,root=d['registration'],years=[2024,2025])
    def test_preview_does_not_create_cases_and_requires_review(self):
        b=self.upload();self.assertEqual(self.client.get('/api/tao/cases').json(),[]);self.assertEqual(self.client.get('/api/cases').json(),[])
        p=self.payload(b);p['confirmed']=False
        self.assertEqual(self.client.post(f"/api/pdf/batches/{b['id']}/apply",json=p,headers=self.h).status_code,422)
    def test_both_modules_once_and_sources_preserved(self):
        b=self.upload();bid=b['id']
        for module in ['tao','kkv']:
            p=self.payload(b,module)
            r=self.client.post(f'/api/pdf/batches/{bid}/apply',json=p,headers=self.h)
            self.assertEqual(r.status_code,200,r.text);cid=r.json()['id']
            base='/api/tao/cases/' if module=='tao' else '/api/cases/'
            self.assertEqual(self.client.get(base+cid).json()['status'],'draft')
            self.assertEqual(len(self.client.get(base+cid+'/documents').json()),1)
            self.assertEqual(self.client.post(f'/api/pdf/batches/{bid}/apply',json=p,headers=self.h).status_code,409)
        self.assertEqual(self.client.get(f'/api/pdf/batches/{bid}/files/'+b['files'][0]['id']).content,b'%PDF-1.4\nsource')
    def test_csrf_client_scope_and_foreign_document(self):
        self.assertEqual(self.client.post('/api/pdf/batches',files={'files':('a.pdf',b'%PDF-1.4')}).status_code,403)
        b=self.upload();p=self.payload(b);p['documents'][0]['document_id']='foreign'
        self.assertEqual(self.client.post(f"/api/pdf/batches/{b['id']}/apply",json=p,headers=self.h).status_code,422)
        self.client.post('/api/users',json=dict(name='Ügyfél',username='client',password=self.password,role='client'),headers=self.h)
        with TestClient(app) as client:
            client.post('/api/auth/login',json=dict(username='client',password=self.password))
            self.assertEqual(client.get(f"/api/pdf/batches/{b['id']}").status_code,403)
    def test_family_sources_and_recheck_preserve_previous_snapshot(self):
        from tests.test_tao_family import assessment
        raw=assessment().model_dump(mode='json')
        created=self.client.post('/api/tao/cases',json={'data':raw},headers=self.h)
        self.assertEqual(created.status_code,200,created.text);cid=created.json()['id']
        uploaded=self.client.post(f'/api/tao/cases/{cid}/documents',files={'file':('nyilatkozat.txt',b'Family evidence','text/plain')},headers=self.h)
        self.assertEqual(uploaded.status_code,200,uploaded.text)
        current=self.client.get('/api/tao/cases/'+cid).json()
        self.assertFalse(current['data']['family_facts'][0]['confirmed'])
        previous=self.client.get(f'/api/tao/cases/{cid}/versions/1').json()
        self.assertTrue(previous['data']['family_facts'][0]['confirmed'])
        family=current['data']['family_facts'][0]
        family['confirmed']=True;family['evidence']['document_id']=uploaded.json()['id']
        updated=self.client.put('/api/tao/cases/'+cid,json={'data':current['data'],'version':current['version']},headers=self.h)
        self.assertEqual(updated.status_code,200,updated.text)
        current=self.client.get('/api/tao/cases/'+cid).json()
        current['data']['family_facts'][0]['evidence']['document_id']='foreign-document'
        rejected=self.client.put('/api/tao/cases/'+cid,json={'data':current['data'],'version':current['version']},headers=self.h)
        self.assertEqual(rejected.status_code,422)

    def test_partial_tao_approval_and_immutable_snapshot(self):
        d=tao().model_dump(mode='json');d.update(scope='A és B, nyilatkozat nélkül',law_date='2024-12-31',law_source='Szakértői jogforrás')
        cid=self.client.post('/api/tao/cases',json={'data':d},headers=self.h).json()['id']
        p=dict(version=1,relationships=True,rules=True,scope=True)
        self.assertEqual(self.client.post(f'/api/tao/cases/{cid}/approve',json=p,headers=self.h).status_code,422)
        p.update(partial=True,note='A rokonsági nyilatkozat hiányzik.')
        r=self.client.post(f'/api/tao/cases/{cid}/approve',json=p,headers=self.h);self.assertEqual(r.status_code,200,r.text)
        report=self.client.get(f'/api/tao/cases/{cid}/report/docx')
        self.assertEqual(report.status_code,200,report.text if report.status_code!=200 else '')
        text='\n'.join(p.text for p in Document(BytesIO(report.content)).paragraphs)
        self.assertIn('RÉSZLEGES',text.upper())
        if shutil.which('soffice'):
            pdf=self.client.get(f'/api/tao/cases/{cid}/report/pdf')
            self.assertEqual(pdf.status_code,200,pdf.text if pdf.status_code!=200 else '')
            self.assertTrue(pdf.content.startswith(b'%PDF'))
        d['title']='=1+2'
        self.client.put('/api/tao/cases/'+cid,json=dict(data=d,version=1),headers=self.h)
        self.assertIsNotNone(self.client.get(f'/api/tao/cases/{cid}/versions/1').json()['approved_at'])
        r=self.client.get(f'/api/tao/cases/{cid}/report/xlsx');self.assertEqual(r.status_code,200,r.text)
        wb=load_workbook(BytesIO(r.content));self.assertFalse(any(c.data_type=='f' for ws in wb for row in ws for c in row))
